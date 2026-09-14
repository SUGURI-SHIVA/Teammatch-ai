from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG_DARK = RGBColor(0x1e, 0x3a, 0x8a)
BG_LIGHT = RGBColor(0xf0, 0xf4, 0xff)
ACCENT = RGBColor(0x25, 0x63, 0xeb)
GREEN = RGBColor(0x16, 0xa3, 0x4a)
WHITE = RGBColor(0xff, 0xff, 0xff)
BLACK = RGBColor(0x1f, 0x29, 0x37)
GRAY = RGBColor(0x6b, 0x72, 0x80)

def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_text(slide, left, top, width, height, text, size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tf

def add_bullet(slide, left, top, width, height, items, size=16, color=BLACK):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
        p.level = 0
    return tf

# Slide 1 - Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG_DARK)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), BG_DARK)
add_shape(slide, Inches(0), Inches(3.2), Inches(13.333), Inches(0.05), ACCENT)
add_text(slide, Inches(1), Inches(1.5), Inches(11), Inches(1.5), "TeamMatch AI", size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(2.8), Inches(11), Inches(1), "AI-Powered Platform for Forming Better Project Teams", size=24, color=RGBColor(0x93, 0xc5, 0xfd), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.8), Inches(11), Inches(0.8), "Connecting Project Ideas with the Right Students", size=20, color=RGBColor(0xbf, 0xdb, 0xfe), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(5.5), Inches(11), Inches(0.8), "Presented by: SUGURI SHIVA", size=18, color=RGBColor(0x93, 0xc5, 0xfd), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(6.2), Inches(11), Inches(0.6), "2026", size=16, color=RGBColor(0x93, 0xc5, 0xfd), align=PP_ALIGN.CENTER)

# Slide 2 - Problem Statement
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "The Problem", size=36, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
add_bullet(slide, Inches(1), Inches(1.8), Inches(11), Inches(5), [
    "Students have project ideas but cannot find the right teammates",
    "Team formation happens through friends, random allocation, or people they already know",
    "Many students have useful skills but don't know which projects need them",
    "No systematic way to match students with projects based on skills and interests",
    "Unbalanced teams with duplicate skills and missing capabilities",
    "Wasted time searching for teammates manually",
], size=20, color=BLACK)

# Slide 3 - Solution
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), GREEN)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "Our Solution: TeamMatch AI", size=36, bold=True, color=WHITE)
add_bullet(slide, Inches(1), Inches(1.8), Inches(5.5), Inches(5), [
    "Direction 1: Project -> Student",
    "Create a project with requirements",
    "AI finds and recommends suitable students",
    "Creator invites students to join",
    "Students accept or decline invitations",
], size=18, color=BLACK)
add_bullet(slide, Inches(7), Inches(1.8), Inches(5.5), Inches(5), [
    "Direction 2: Student -> Project",
    "Build a skills and interests profile",
    "AI discovers matching projects",
    "Student sees match score and reasons",
    "Student requests to join a project",
], size=18, color=BLACK)
add_shape(slide, Inches(4.5), Inches(1.5), Inches(4.3), Inches(5.2), RGBColor(0xf0, 0xfd, 0xf4))
add_text(slide, Inches(4.7), Inches(1.8), Inches(4), Inches(0.6), "Two-Way Matching", size=20, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
add_text(slide, Inches(4.7), Inches(2.5), Inches(4), Inches(3.5), "The core innovation: AI creates connections in BOTH directions between project ideas and students who want to build them.", size=16, color=BLACK, align=PP_ALIGN.CENTER)

# Slide 4 - Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "System Architecture", size=36, bold=True, color=WHITE)

boxes = [
    (0.5, 1.8, 3, 1.2, "React + TypeScript\nFrontend", RGBColor(0xdb, 0xea, 0xfe)),
    (4.5, 1.8, 3, 1.2, "Tailwind CSS\nStyling", RGBColor(0xf3, 0xe8, 0xff)),
    (8.5, 1.8, 3, 1.2, "Vite\nBuild Tool", RGBColor(0xec, 0xfd, 0xf5)),
    (2.5, 3.5, 3, 1.2, "Express.js\nBackend API", RGBColor(0xfe, 0xf3, 0xc7)),
    (7.5, 3.5, 3, 1.2, "Prisma ORM\nDatabase Layer", RGBColor(0xff, 0xed, 0xd5)),
    (1.5, 5.2, 3, 1.2, "SQLite\nDatabase", RGBColor(0xdc, 0xfc, 0xe7)),
    (5.5, 5.2, 3, 1.2, "JWT Auth\nAuthentication", RGBColor(0xe0, 0xf2, 0xfe)),
    (9.5, 5.2, 3, 1.2, "AI Matching\nEngine", RGBColor(0xfd, 0xe6, 0x8a)),
]
for x, y, w, h, text, color in boxes:
    shape = add_shape(slide, Inches(x), Inches(y), Inches(w), Inches(h), color)
    shape.line.color.rgb = RGBColor(0xd1, 0xd5, 0xdb)
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(12)

# Slide 5 - AI Matching Engine
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "AI Matching Engine", size=36, bold=True, color=WHITE)
add_text(slide, Inches(1), Inches(1.5), Inches(11), Inches(0.6), "Weighted Scoring Algorithm", size=22, bold=True, color=BLACK)

scoring = [
    ("Skill Match", "40%", "Compares project required skills with student skills"),
    ("Interest Match", "20%", "Matches project domain interests with student interests"),
    ("Role Match", "15%", "Checks if student's preferred role matches project needs"),
    ("Experience Match", "10%", "Compares experience level (Beginner/Intermediate/Advanced)"),
    ("Availability Match", "10%", "Checks overlapping available days/hours"),
    ("Learning Goals", "5%", "Bonus for skills student wants to learn"),
]
for i, (name, weight, desc) in enumerate(scoring):
    y = 2.3 + i * 0.7
    add_shape(slide, Inches(1), Inches(y), Inches(2), Inches(0.5), ACCENT if i % 2 == 0 else RGBColor(0x1e, 0x40, 0xaf))
    add_text(slide, Inches(1.1), Inches(y + 0.05), Inches(1.8), Inches(0.4), f"{name}: {weight}", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(3.5), Inches(y + 0.05), Inches(9), Inches(0.4), desc, size=14, color=BLACK)

# Slide 6 - Key Features
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "Key Features", size=36, bold=True, color=WHITE)

features = [
    ("AI Recommendations", "Smart matching with explainable scores and reasons"),
    ("Two-Way Matching", "Project finds students AND students find projects"),
    ("Team Coverage", "Analyzes skill coverage, missing skills, duplicates"),
    ("Invitation System", "Project creators invite, students accept/decline"),
    ("Join Requests", "Students request, creators accept/reject"),
    ("Notifications", "Real-time alerts for all team activities"),
    ("Search & Filters", "Filter by skill, role, experience, domain"),
    ("Theme Customizer", "6 color themes users can switch between"),
    ("Responsive Design", "Works on desktop, tablet, and mobile"),
]
for i, (title, desc) in enumerate(features):
    col = i % 3
    row = i // 3
    x = 0.5 + col * 4.2
    y = 1.6 + row * 1.8
    shape = add_shape(slide, Inches(x), Inches(y), Inches(3.8), Inches(1.5), RGBColor(0xf0, 0xf4, 0xff))
    shape.line.color.rgb = RGBColor(0xdb, 0xea, 0xfe)
    shape.line.width = Pt(1)
    add_text(slide, Inches(x + 0.2), Inches(y + 0.2), Inches(3.4), Inches(0.5), title, size=16, bold=True, color=ACCENT)
    add_text(slide, Inches(x + 0.2), Inches(y + 0.7), Inches(3.4), Inches(0.6), desc, size=13, color=GRAY)

# Slide 7 - User Flow
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "User Flow", size=36, bold=True, color=WHITE)

flow_items = [
    "Register / Login",
    "Create Profile or\nCreate Project",
    "AI Analyzes\nRequirements",
    "AI Matches\nStudents/Projects",
    "View Match %\n& Reasons",
    "Invite or\nRequest to Join",
    "Accept / Decline",
    "Team Formed!"
]
for i, text in enumerate(flow_items):
    x = 0.3 + i * 1.6
    color = ACCENT if i < 4 else GREEN
    shape = add_shape(slide, Inches(x), Inches(2), Inches(1.3), Inches(1.3), color)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(8)
    if i < len(flow_items) - 1:
        add_text(slide, Inches(x + 1.3), Inches(2.3), Inches(0.3), Inches(0.5), "->", size=20, bold=True, color=GRAY, align=PP_ALIGN.CENTER)

add_text(slide, Inches(0.5), Inches(4), Inches(12), Inches(0.6), "Project Creator Flow", size=20, bold=True, color=ACCENT)
add_bullet(slide, Inches(0.5), Inches(4.6), Inches(5.5), Inches(2.5), [
    "1. Create project with requirements",
    "2. AI recommends suitable students",
    "3. View match scores and reasons",
    "4. Invite students",
    "5. Students accept or decline",
    "6. Review join requests",
], size=14, color=BLACK)

add_text(slide, Inches(7), Inches(4), Inches(5.5), Inches(0.6), "Student Flow", size=20, bold=True, color=GREEN)
add_bullet(slide, Inches(7), Inches(4.6), Inches(5.5), Inches(2.5), [
    "1. Build skills and interests profile",
    "2. AI recommends matching projects",
    "3. View match scores and reasons",
    "4. Request to join projects",
    "5. Creator reviews and accepts",
    "6. Join the team",
], size=14, color=BLACK)

# Slide 8 - Database Schema
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "Database Schema", size=36, bold=True, color=WHITE)

tables = [
    ("users", "id, email, password,\nname, role"),
    ("student_profiles", "userId, college, dept,\nyear, experience"),
    ("projects", "id, title, description,\ndomain, teamSize"),
    ("skills", "id, name, category"),
    ("interests", "id, name"),
    ("roles", "id, name"),
    ("project_members", "projectId, userId, role"),
    ("invitations", "projectId, senderId,\nreceiverId, status"),
    ("join_requests", "projectId, studentId,\nstatus, matchScore"),
    ("notifications", "userId, type, title,\nmessage, read"),
]
for i, (table, fields) in enumerate(tables):
    col = i % 5
    row = i // 5
    x = 0.3 + col * 2.6
    y = 1.5 + row * 2.8
    shape = add_shape(slide, Inches(x), Inches(y), Inches(2.3), Inches(2.5), RGBColor(0xf0, 0xf4, 0xff))
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(1)
    add_text(slide, Inches(x + 0.1), Inches(y + 0.1), Inches(2.1), Inches(0.5), table, size=13, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(x + 0.1), Inches(y + 0.7), Inches(2.1), Inches(1.5), fields, size=11, color=GRAY, align=PP_ALIGN.CENTER)

# Slide 9 - Tech Stack
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "Technology Stack", size=36, bold=True, color=WHITE)

stack = [
    ("Frontend", ["React 18", "TypeScript", "Tailwind CSS", "Vite", "React Router", "Lucide Icons"], RGBColor(0xdb, 0xea, 0xfe)),
    ("Backend", ["Node.js", "Express.js", "TypeScript", "Prisma ORM", "JWT Auth", "bcryptjs"], RGBColor(0xec, 0xfd, 0xf5)),
    ("Database", ["SQLite", "Prisma Client", "14 Tables", "UUID Primary Keys", "Foreign Keys", "Unique Constraints"], RGBColor(0xf3, 0xe8, 0xff)),
    ("AI Engine", ["Weighted Scoring", "Skill Aliases", "Cosine Similarity", "Team Coverage", "Role Matching", "Availability Check"], RGBColor(0xfd, 0xe6, 0x8a)),
]
for i, (title, items, color) in enumerate(stack):
    x = 0.3 + i * 3.2
    shape = add_shape(slide, Inches(x), Inches(1.5), Inches(3), Inches(5.5), color)
    shape.line.fill.background()
    add_text(slide, Inches(x + 0.2), Inches(1.7), Inches(2.6), Inches(0.5), title, size=20, bold=True, color=BLACK, align=PP_ALIGN.CENTER)
    add_bullet(slide, Inches(x + 0.3), Inches(2.4), Inches(2.5), Inches(4), items, size=15, color=BLACK)

# Slide 10 - Deployment
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "Deployment & Access", size=36, bold=True, color=WHITE)

deploy_items = [
    ("GitHub", "Source code version control\nhttps://github.com/SUGURI-SHIVA/Teammatch-ai", RGBColor(0x24, 0x29, 0x2e)),
    ("Render.com", "Free cloud hosting\nAuto-deploy from GitHub\nhttps://teammatch-ai.onrender.com", RGBColor(0x66, 0x33, 0x99)),
    ("SQLite", "Lightweight database\nNo external DB needed\nPortable and simple", RGBColor(0x00, 0x3b, 0x57)),
]
for i, (title, desc, color) in enumerate(deploy_items):
    x = 0.5 + i * 4.2
    shape = add_shape(slide, Inches(x), Inches(1.8), Inches(3.8), Inches(2.5), color)
    shape.line.fill.background()
    add_text(slide, Inches(x + 0.2), Inches(2), Inches(3.4), Inches(0.6), title, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(x + 0.2), Inches(2.7), Inches(3.4), Inches(1.5), desc, size=14, color=WHITE, align=PP_ALIGN.CENTER)

add_text(slide, Inches(1), Inches(4.8), Inches(11), Inches(0.6), "Access Methods", size=24, bold=True, color=BLACK)
add_bullet(slide, Inches(1), Inches(5.5), Inches(11), Inches(2), [
    "Online: https://teammatch-ai.onrender.com (any device, anywhere)",
    "Local: http://localhost:5173 (on your PC)",
    "Custom URL: http://teammatchai.local (on your PC)",
    "Phone: Add to home screen from browser menu",
], size=16, color=BLACK)

# Slide 11 - Demo
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "Demo", size=36, bold=True, color=WHITE)

add_text(slide, Inches(1), Inches(1.8), Inches(11), Inches(0.6), "Live at: https://teammatch-ai.onrender.com", size=24, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

add_text(slide, Inches(1), Inches(2.8), Inches(11), Inches(0.6), "Demo Accounts (Password: password123)", size=20, bold=True, color=BLACK, align=PP_ALIGN.CENTER)

accounts = [
    ("alice@college.edu", "Alice Johnson", "Both"),
    ("bob@college.edu", "Bob Smith", "Student"),
    ("carol@college.edu", "Carol Davis", "Student"),
    ("dave@college.edu", "Dave Wilson", "Student"),
]
for i, (email, name, role) in enumerate(accounts):
    y = 3.5 + i * 0.6
    add_text(slide, Inches(2), Inches(y), Inches(3), Inches(0.5), email, size=16, color=BLACK, align=PP_ALIGN.LEFT)
    add_text(slide, Inches(5.5), Inches(y), Inches(3), Inches(0.5), name, size=16, color=BLACK, align=PP_ALIGN.LEFT)
    add_text(slide, Inches(8.5), Inches(y), Inches(2), Inches(0.5), role, size=16, bold=True, color=ACCENT, align=PP_ALIGN.LEFT)

# Slide 12 - Future Scope
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), ACCENT)
add_text(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.8), "Future Scope", size=36, bold=True, color=WHITE)

future = [
    ("LLM Integration", "Use GPT/Claude for deeper\nproject requirement analysis"),
    ("Real-time Chat", "Built-in messaging between\nteam members"),
    ("Mobile App", "Native Android/iOS app\nusing React Native"),
    ("Advanced AI", "Deep learning models for\nbetter match predictions"),
    ("Video Profiles", "Students can add video\nintros to their profiles"),
    ("Analytics Dashboard", "Team performance tracking\nand project insights"),
]
for i, (title, desc) in enumerate(future):
    col = i % 3
    row = i // 3
    x = 0.5 + col * 4.2
    y = 1.6 + row * 2.8
    shape = add_shape(slide, Inches(x), Inches(y), Inches(3.8), Inches(2.3), RGBColor(0xf0, 0xf4, 0xff))
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(1)
    add_text(slide, Inches(x + 0.2), Inches(y + 0.3), Inches(3.4), Inches(0.5), title, size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(x + 0.2), Inches(y + 1), Inches(3.4), Inches(1), desc, size=14, color=BLACK, align=PP_ALIGN.CENTER)

# Slide 13 - Thank You
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG_DARK)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), BG_DARK)
add_text(slide, Inches(1), Inches(2.5), Inches(11), Inches(1.5), "Thank You!", size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(4.2), Inches(11), Inches(0.8), "TeamMatch AI — Connecting Ideas with Talent", size=24, color=RGBColor(0x93, 0xc5, 0xfd), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(5.5), Inches(11), Inches(0.6), "https://teammatch-ai.onrender.com", size=18, color=RGBColor(0xbf, 0xdb, 0xfe), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(6.2), Inches(11), Inches(0.6), "Questions?", size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

prs.save(r"C:\Users\sugur\OneDrive\Desktop\shiva f\TeamMatch_AI_Presentation.pptx")
print("Presentation saved!")
